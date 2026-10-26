"""CSV/XLSX export of archives and statistics (services/export.py)."""

import csv
import io
from datetime import datetime, timedelta, timezone

import pytest
from httpx import AsyncClient
from openpyxl import load_workbook

pytestmark = pytest.mark.integration


def _csv(response) -> list[list[str]]:
    return list(csv.reader(io.StringIO(response.content.decode("utf-8"))))


@pytest.fixture
async def archives(printer_factory, archive_factory, db_session):
    from backend.app.models.project import Project

    project = Project(name="Gift boxes")
    db_session.add(project)
    await db_session.commit()
    await db_session.refresh(project)

    printer = await printer_factory()
    other = await printer_factory(name="Other")
    now = datetime.now(timezone.utc)
    a = await archive_factory(
        printer.id,
        print_name="Gear box",
        status="completed",
        project_id=project.id,
        designer="Ada",
        started_at=now - timedelta(hours=2),
        completed_at=now - timedelta(hours=1),
    )
    b = await archive_factory(printer.id, print_name="Clip", status="failed", failure_reason="Clogged nozzle")
    c = await archive_factory(other.id, print_name="Hinge", status="completed", notes="for the gearbox lid")
    return {"printer": printer, "other": other, "project": project, "gear": a, "clip": b, "hinge": c}


class TestArchiveExport:
    async def test_csv_with_default_fields(self, async_client: AsyncClient, archives):
        response = await async_client.get("/api/v1/archives/export")

        assert response.status_code == 200
        assert response.headers["content-type"].startswith("text/csv")
        assert "archives_export_" in response.headers["content-disposition"]
        header, *rows = _csv(response)
        assert header[:4] == ["ID", "Print Name", "Filename", "Status"]
        by_name = {row[1]: dict(zip(header, row, strict=True)) for row in rows}
        assert set(by_name) == {"Gear box", "Clip", "Hinge"}
        gear = by_name["Gear box"]
        assert (gear["Project"], gear["Designer"]) == ("Gift boxes", "Ada")
        assert gear["Started At"].startswith(str(datetime.now(timezone.utc).year))
        assert by_name["Clip"]["Failure Reason"] == "Clogged nozzle"
        assert by_name["Hinge"]["Project"] == ""

    async def test_selected_fields_and_filters(self, async_client: AsyncClient, archives):
        printer_id = archives["printer"].id

        response = await async_client.get(
            "/api/v1/archives/export",
            params={"fields": "print_name, status", "printer_id": printer_id, "status": "completed"},
        )

        assert _csv(response) == [["Print Name", "Status"], ["Gear box", "completed"]]

    async def test_project_and_search_filters(self, async_client: AsyncClient, archives):
        by_project = await async_client.get(
            "/api/v1/archives/export", params={"fields": "print_name", "project_id": archives["project"].id}
        )
        by_search = await async_client.get("/api/v1/archives/export", params={"fields": "print_name", "search": "gear"})

        assert _csv(by_project)[1:] == [["Gear box"]]
        # Search covers names, tags, notes and designer.
        assert sorted(r[0] for r in _csv(by_search)[1:]) == ["Gear box", "Hinge"]

    async def test_date_range(self, async_client: AsyncClient, archives):
        future = (datetime.now() + timedelta(days=1)).isoformat()
        past = (datetime.now() - timedelta(days=1)).isoformat()

        none = await async_client.get("/api/v1/archives/export", params={"fields": "id", "date_from": future})
        everything = await async_client.get(
            "/api/v1/archives/export", params={"fields": "id", "date_from": past, "date_to": future}
        )

        assert len(_csv(none)) == 1
        assert len(_csv(everything)) == 4

    async def test_xlsx(self, async_client: AsyncClient, archives):
        response = await async_client.get(
            "/api/v1/archives/export", params={"format": "xlsx", "fields": "print_name,cost"}
        )

        assert response.headers["content-type"].startswith("application/vnd.openxmlformats")
        sheet = load_workbook(io.BytesIO(response.content)).active
        assert sheet.title == "Archives"
        assert [c.value for c in sheet[1]] == ["Print Name", "Cost"]
        assert sheet[1][0].font.bold is True
        assert sheet.freeze_panes == "A2"
        assert {sheet.cell(row=r, column=1).value for r in range(2, 5)} == {"Gear box", "Clip", "Hinge"}

    @pytest.mark.parametrize(
        "params",
        [{"format": "pdf"}, {"date_from": "yesterday"}, {"date_to": "31/12/2026"}],
    )
    async def test_bad_requests(self, async_client: AsyncClient, archives, params):
        response = await async_client.get("/api/v1/archives/export", params=params)
        assert response.status_code == 400


class TestStatsExport:
    async def test_csv_summary(self, async_client: AsyncClient, archives):
        response = await async_client.get("/api/v1/archives/stats/export", params={"days": 7})

        assert response.status_code == 200
        rows = _csv(response)
        assert rows[0] == ["Metric", "Value"]
        metrics = {row[0]: row[1] for row in rows if len(row) == 2}
        assert metrics["Period (days)"] == "7"
        assert "Total Prints" in metrics
        sections = [row[0] for row in rows if len(row) == 2 and row[1] == ""]
        assert sections == ["Failures by Reason", "Failures by Filament", "Failures by Printer", "Weekly Trend"]

    async def test_xlsx_summary(self, async_client: AsyncClient, archives):
        response = await async_client.get("/api/v1/archives/stats/export", params={"format": "xlsx"})

        sheet = load_workbook(io.BytesIO(response.content)).active
        assert sheet.title == "Statistics"
        assert [c.value for c in sheet[1]][:2] == ["Metric", "Value"]

    async def test_bad_format(self, async_client: AsyncClient, archives):
        response = await async_client.get("/api/v1/archives/stats/export", params={"format": "pdf"})
        assert response.status_code == 400
